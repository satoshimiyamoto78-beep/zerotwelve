"""Build the contractor BOQ workbook (EN + VI) with product images, swatches and input columns."""
import os, sys
from PIL import Image as PI
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XI
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
HERE = os.path.dirname(os.path.abspath(__file__)); S = os.path.dirname(HERE)
sys.path.insert(0, HERE)
from boq_data import SECTIONS, MATERIALS, CHECKS
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, 'Samarkand_BOQ_R12-v8.2_EN-VI.xlsx')
SHOTS, TEX, THUMB = os.path.join(HERE, 'shots'), os.path.join(S, 'r3d', 'tex2'), os.path.join(HERE, 'thumbs')
REFS = '/home/user/zerotwelve/samarkand/refs'
os.makedirs(THUMB, exist_ok=True)
VIEWER = 'https://claude.ai/artifact/7NGDiBpaXwuSr9gytuvAtm'
SW = {'sw_marble': 'marblegrid.png', 'sw_zcarpet': 'zcarpet.png', 'sw_alley': 'alleyrun.png',
      'sw_rug_persian': 'rug_oldpersian.jpg', 'sw_suzani': 'suzani_w.png'}
# drawing references corrected to the sheet titles of the R12-v8 set
REF_FIX = {'E05': 'S-06 / 3D', 'E06': 'S-06 / 3D', 'E07': '3D viewer', 'E09': '3D viewer', 'F01': '3D viewer', 'F02': '3D viewer',
           'F03': '3D viewer', 'J01': 'S-07', 'J02': 'S-07', 'J03': 'S-07', 'E04': 'Plan / 3D'}
IMG_FIX = {'E04': None}

F = 'Arial'
def font(**k): return Font(name=F, **k)
NAVY, GOLD = '1F3A5F', 'C9A227'
fill = lambda c: PatternFill('solid', start_color=c, end_color=c)
YEL, SEC, HEAD, GREY = fill('FFF2CC'), fill('E8EEF5'), fill(NAVY), fill('F3F3F3')
thin = Side(style='thin', color='BFBFBF'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical='top'); CEN = Alignment(horizontal='center', vertical='center', wrap_text=True)
VND = '#,##0;(#,##0);"-"'

def thumb(key, w, h, crop_square=False):
    if not key: return None
    src = os.path.join(TEX, SW[key]) if key.startswith('sw_') else (key if os.path.isabs(key) else os.path.join(SHOTS, key + '.png'))
    if not os.path.exists(src): print('no image', key); return None
    im = PI.open(src).convert('RGB')
    if crop_square or key.startswith('sw_'):
        s = min(im.size) // (2 if key.startswith('sw_') or crop_square else 1); cx, cy = im.size[0] // 2, im.size[1] // 2
        im = im.crop((cx - s // 2, cy - s // 2, cx + s // 2, cy + s // 2))
    else:  # trim the studio background around the object
        g = im.convert('L').point(lambda v: 255 if v < 215 else 0); bb = g.getbbox()
        if bb:
            pad = 12; bb = (max(0, bb[0] - pad), max(0, bb[1] - pad), min(im.size[0], bb[2] + pad), min(im.size[1], bb[3] + pad)); im = im.crop(bb)
    im.thumbnail((w * 2, h * 2), PI.LANCZOS)
    canvas = PI.new('RGB', (w * 2, h * 2), (255, 255, 255)); canvas.paste(im, ((w * 2 - im.size[0]) // 2, (h * 2 - im.size[1]) // 2))
    p = os.path.join(THUMB, os.path.basename(src).rsplit('.', 1)[0] + f'_{w}x{h}.png'); canvas.save(p)
    x = XI(p); x.width, x.height = w, h; return x

wb = Workbook()
# ---------------------------------------------------------------- Read me
rd = wb.active; rd.title = 'Read me'
rd.sheet_view.showGridLines = False
for c, w in zip('ABCD', (3, 34, 62, 62)): rd.column_dimensions[c].width = w
r = 2
rd.cell(r, 2, 'SAMARKAND — terrace restaurant, Da Nang').font = font(bold=True, size=18, color=NAVY); r += 1
rd.cell(r, 2, 'Bill of Quantities for pricing / Bảng khối lượng để báo giá — design revision R12-v8.2').font = font(size=12, color='595959'); r += 2
info = [
 ('Project / Dự án', 'Tajik & Uzbek restaurant terrace: new steel roof with a glazed-tile dome, aivan platform, themed booths, tents, tapchan, bar, mangal grill, street facade.',
  'Sân thượng nhà hàng Tajik & Uzbek: mái thép mới với mái vòm gạch men, bục aivan, ô ngồi theo chủ đề, lều, tapchan, quầy bar, bếp nướng, mặt tiền.'),
 ('Size / Quy mô', 'Terrace ≈ 317 m² + kids 17 m² + mangal 19 m² · roof ≈ 413 m² · 155 seats (+31 street bench seats)',
  'Sân ≈ 317 m² + trẻ em 17 m² + bếp nướng 19 m² · mái ≈ 413 m² · 155 chỗ ngồi (+31 chỗ ghế băng mặt đường)'),
 ('What we ask / Yêu cầu', 'Please fill the YELLOW cells on sheet "BOQ": unit price in VND (supply + install), your note, and a status for every line. Totals calculate automatically on "Summary".',
  'Vui lòng điền các ô MÀU VÀNG trong sheet "BOQ": đơn giá VND (cung cấp + lắp đặt), ghi chú, và trạng thái cho từng dòng. Tổng tiền tự tính ở sheet "Summary".'),
 ('Status values / Trạng thái', 'Priced = price given · Need info = you need a drawing / sample / answer · Alternative = you propose another material or method (describe in Note) · Not in scope = you do not do this item',
  'Priced = đã báo giá · Need info = cần thêm bản vẽ / mẫu / trả lời · Alternative = đề xuất vật liệu hoặc cách làm khác (ghi trong Note) · Not in scope = không làm hạng mục này'),
 ('Quantities / Khối lượng', 'Taken from our CAD plan and 3D model. They are for comparison of offers — please check on site and correct in your Note if you measure differently.',
  'Lấy từ bản vẽ CAD và mô hình 3D. Dùng để so sánh báo giá — vui lòng kiểm tra tại công trình và ghi chú nếu đo khác.'),
 ('Images / Hình ảnh', 'Each line has a picture from the 3D model. Sheet "Materials" shows the finishes (codes M01–M30), sheet "References" shows real photos the owner likes.',
  'Mỗi dòng có hình từ mô hình 3D. Sheet "Materials" là mẫu vật liệu hoàn thiện (mã M01–M30), sheet "References" là ảnh thực tế chủ đầu tư thích.'),
 ('Please check first / Kiểm tra trước', 'Sheet "Design checks" lists 15 technical points (structure, fire, slip, gas, permits). Please read before pricing and comment.',
  'Sheet "Design checks" liệt kê 15 điểm kỹ thuật (kết cấu, PCCC, chống trơn, gas, giấy phép). Vui lòng đọc trước khi báo giá và góp ý.'),
 ('3D model / Mô hình 3D', VIEWER + '  (open in a browser; buttons: roof, ceiling, facade, dimensions)', VIEWER + '  (mở bằng trình duyệt; nút: mái, trần, mặt tiền, kích thước)'),
 ('Drawings / Bản vẽ', 'Plan R12-v8 (EN/RU) · Roof plan · S-01 Roof framing · S-02 Central dome · S-03 Chinar + skylight · S-04 Majolica column · S-05 Street portal (pishtaq) + sebil · S-06 Khiva arcade · S-07 Bar · S-08 Themed booth',
  'Mặt bằng R12-v8 · Mặt bằng mái · S-01 Khung mái · S-02 Mái vòm · S-03 Cây chinar + giếng trời · S-04 Cột majolica · S-05 Cổng pishtaq + sebil · S-06 Hàng vòm Khiva · S-07 Quầy bar · S-08 Ô ngồi chủ đề'),
 ('Prices / Giá', 'VND, excluding VAT, including materials, labour, transport, installation, scaffolding and site cleaning. Give a separate line if something is excluded.',
  'VND, chưa VAT, bao gồm vật tư, nhân công, vận chuyển, lắp đặt, giàn giáo và vệ sinh. Ghi riêng nếu có phần không bao gồm.'),
]
for j, t in enumerate(('', 'English', 'Tiếng Việt')):
    c = rd.cell(r, 2 + j, t); c.font = font(bold=True, color='FFFFFF'); c.fill = HEAD
r += 1
for k, en, vi in info:
    rd.cell(r, 2, k).font = font(bold=True, color=NAVY)
    rd.cell(r, 3, en); rd.cell(r, 4, vi)
    for j in (2, 3, 4): rd.cell(r, j).alignment = WRAP; rd.cell(r, j).border = BOX
    if k.startswith('3D'): rd.cell(r, 3).hyperlink = VIEWER; rd.cell(r, 3).font = font(color='0563C1', underline='single')
    rd.row_dimensions[r].height = 48; r += 1
r += 1
rd.cell(r, 2, 'Example of a filled line / Ví dụ một dòng đã điền').font = font(bold=True, color=NAVY); r += 1
ex_h = ('Code', 'Item', 'Unit', 'Qty', 'Unit price VND', 'Amount VND', 'Note', 'Status')
ex_v = ('F01', 'Carved armchair', 'pcs', 46, 4800000, None, 'Acacia wood, walnut stain; prototype 10 days', 'Priced')
# compact example table under the info block
r0 = r
for j, t in enumerate(ex_h):
    c = rd.cell(r0 + 1 + j, 2, t); c.font = font(bold=True); c.border = BOX; c.fill = GREY
    v = ex_v[j]; c2 = rd.cell(r0 + 1 + j, 3, v if j != 5 else f'=C{r0 + 4}*C{r0 + 5}'); c2.border = BOX
    if j in (4, 5): c2.number_format = VND
    if j in (4, 6, 7): c2.fill = YEL
rd.cell(r0 + 9, 2, 'Yellow cells = your input (ô vàng = nhà thầu điền). Everything else is formula or our data.').font = font(italic=True, color='7F6000')

# ---------------------------------------------------------------- BOQ
ws = wb.create_sheet('BOQ')
cols = [('Code\nMã', 7), ('Item (EN)', 34), ('Hạng mục (VI)', 34), ('Specification / Quy cách', 34), ('Size / Kích thước', 14),
        ('Unit\nĐV', 6), ('Qty\nSL', 9), ('Image / Hình', 25), ('Ref\nBản vẽ', 10),
        ('Unit price VND\nĐơn giá', 15), ('Amount VND\nThành tiền', 17), ('Contractor note\nGhi chú nhà thầu', 30), ('Status\nTrạng thái', 13)]
for j, (t, w) in enumerate(cols, 1):
    c = ws.cell(1, j, t); c.font = font(bold=True, color='FFFFFF', size=10); c.fill = HEAD; c.alignment = CEN; c.border = BOX
    ws.column_dimensions[c.column_letter].width = w
ws.row_dimensions[1].height = 34; ws.freeze_panes = 'B2'
dv = DataValidation(type='list', formula1='"Priced,Need info,Alternative,Not in scope"', allow_blank=True)
dv.error = 'Choose: Priced / Need info / Alternative / Not in scope'; ws.add_data_validation(dv)
row = 2; sec_rows = []
for code, en, vi, items in SECTIONS:
    hr = row
    ws.cell(hr, 1, code); ws.cell(hr, 2, en); ws.cell(hr, 3, vi)
    for j in range(1, 14):
        c = ws.cell(hr, j); c.fill = SEC; c.font = font(bold=True, color=NAVY, size=11); c.border = BOX; c.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[hr].height = 24; row += 1; first = row
    for it in items:
        ic, ien, ivi, spec, size, unit, qty, img, ref = it
        ref = REF_FIX.get(ic, ref); img = IMG_FIX.get(ic, img)
        vals = (ic, ien, ivi, spec, size, unit, qty, None, ref)
        for j, v in enumerate(vals, 1):
            c = ws.cell(row, j, v); c.font = font(size=9); c.alignment = WRAP; c.border = BOX
        ws.cell(row, 7).number_format = '#,##0.0' if isinstance(qty, float) else '#,##0'
        ws.cell(row, 7).font = font(size=9, color='0000FF')
        ws.cell(row, 1).font = font(size=9, bold=True)
        for j in (10, 12, 13):
            c = ws.cell(row, j); c.fill = YEL; c.border = BOX; c.font = font(size=9); c.alignment = WRAP
        ws.cell(row, 10).number_format = VND
        a = ws.cell(row, 11, f'=IF(J{row}="","",G{row}*J{row})'); a.number_format = VND; a.font = font(size=9); a.border = BOX; a.alignment = Alignment(vertical='top')
        dv.add(ws.cell(row, 13))
        x = thumb(img, 170, 120)
        if x: ws.add_image(x, f'H{row}'); ws.row_dimensions[row].height = 96
        else: ws.row_dimensions[row].height = 60
        row += 1
    last = row - 1
    s = ws.cell(hr, 11, f'=SUM(K{first}:K{last})'); s.number_format = VND
    ws.cell(hr, 10, 'Subtotal').alignment = Alignment(horizontal='right', vertical='center')
    sec_rows.append((code, en, vi, hr, first, last))
    row += 0
ws.cell(1, 7).comment = Comment('Blue = quantity from our CAD / 3D model. Check on site; write corrections in the Note column.', 'Designer')
ws.auto_filter.ref = f'A1:M{row - 1}'
ws.page_setup.orientation = 'landscape'; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.print_title_rows = '1:1'
nboq = row - 1

# ---------------------------------------------------------------- Summary
sm = wb.create_sheet('Summary', 1)
sm.sheet_view.showGridLines = False
for c, w in zip('ABCDEF', (3, 8, 44, 44, 20, 14)): sm.column_dimensions[c].width = w
sm.cell(2, 2, 'Summary / Tổng hợp (VND, excl. VAT unless stated)').font = font(bold=True, size=14, color=NAVY)
hdr = ('Code', 'Section', 'Hạng mục', 'Amount VND', 'Lines priced')
for j, t in enumerate(hdr):
    c = sm.cell(4, 2 + j, t); c.font = font(bold=True, color='FFFFFF'); c.fill = HEAD; c.border = BOX; c.alignment = CEN
r = 5
for code, en, vi, hr, first, last in sec_rows:
    sm.cell(r, 2, code); sm.cell(r, 3, en); sm.cell(r, 4, vi)
    c = sm.cell(r, 5, f'=BOQ!K{hr}'); c.number_format = VND; c.font = font(color='008000')
    sm.cell(r, 6, f'=COUNTIF(BOQ!M{first}:M{last},"Priced")&" / {last - first + 1}"').alignment = CEN
    for j in range(2, 7): sm.cell(r, j).border = BOX; sm.cell(r, j).font = font(color='008000' if j == 5 else '000000')
    r += 1
tr = r
sm.cell(tr, 3, 'Subtotal works / Tổng cộng').font = font(bold=True)
c = sm.cell(tr, 5, f'=SUM(E5:E{tr - 1})'); c.number_format = VND; c.font = font(bold=True)
sm.cell(tr + 1, 3, 'Contingency / Dự phòng (%)'); c = sm.cell(tr + 1, 6, 0.07); c.number_format = '0%'; c.fill = YEL; c.font = font(color='0000FF')
c = sm.cell(tr + 1, 5, f'=E{tr}*F{tr + 1}'); c.number_format = VND
sm.cell(tr + 2, 3, 'Total before VAT / Tổng trước VAT').font = font(bold=True)
c = sm.cell(tr + 2, 5, f'=E{tr}+E{tr + 1}'); c.number_format = VND; c.font = font(bold=True)
sm.cell(tr + 3, 3, 'VAT (%)'); c = sm.cell(tr + 3, 6, 0.08); c.number_format = '0%'; c.fill = YEL; c.font = font(color='0000FF')
c = sm.cell(tr + 3, 5, f'=E{tr + 2}*F{tr + 3}'); c.number_format = VND
sm.cell(tr + 4, 3, 'GRAND TOTAL / TỔNG CỘNG').font = font(bold=True, size=12, color=NAVY)
c = sm.cell(tr + 4, 5, f'=E{tr + 2}+E{tr + 3}'); c.number_format = VND; c.font = font(bold=True, size=12, color=NAVY)
for rr in range(tr, tr + 5):
    for j in range(2, 7): sm.cell(rr, j).border = BOX
sm.cell(tr + 6, 3, 'Lines with status "Need info"').font = font(italic=True)
sm.cell(tr + 6, 5, f'=COUNTIF(BOQ!M2:M{nboq},"Need info")')
sm.cell(tr + 7, 3, 'Lines with status "Alternative"').font = font(italic=True)
sm.cell(tr + 7, 5, f'=COUNTIF(BOQ!M2:M{nboq},"Alternative")')
sm.cell(tr + 8, 3, 'Lines without status yet').font = font(italic=True)
sm.cell(tr + 8, 5, f'=COUNTA(BOQ!A2:A{nboq})-{len(sec_rows)}-COUNTA(BOQ!M2:M{nboq})')
sm.cell(tr + 10, 3, 'Contingency 7 % and VAT 8 % are editable defaults (yellow). VAT rate per current Vietnamese rules — adjust if your invoice uses 10 %.').font = font(italic=True, size=9, color='7F6000')

# ---------------------------------------------------------------- Materials
mt = wb.create_sheet('Materials')
mcols = (('Code', 7), ('Sample / Mẫu', 20), ('Material (EN)', 30), ('Vật liệu (VI)', 30), ('Where used / Vị trí', 36), ('Spec / Quy cách', 36), ('Contractor note / Ghi chú', 30))
for j, (t, w) in enumerate(mcols, 1):
    c = mt.cell(1, j, t); c.font = font(bold=True, color='FFFFFF'); c.fill = HEAD; c.alignment = CEN; c.border = BOX
    mt.column_dimensions[c.column_letter].width = w
mt.freeze_panes = 'A2'
for i, (code, en, vi, where, spec, f_) in enumerate(MATERIALS, 2):
    for j, v in enumerate((code, None, en, vi, where, spec, None), 1):
        c = mt.cell(i, j, v); c.font = font(size=9, bold=(j == 1)); c.alignment = WRAP; c.border = BOX
    mt.cell(i, 7).fill = YEL
    p = os.path.join(TEX, f_)
    if os.path.exists(p):
        im = PI.open(p).convert('RGB'); s = min(im.size) // 2; cx, cy = im.size[0] // 2, im.size[1] // 2
        im = im.crop((cx - s // 2, cy - s // 2, cx + s // 2, cy + s // 2)).resize((220, 220), PI.LANCZOS)
        tp = os.path.join(THUMB, 'mat_' + code + '.png'); im.save(tp); x = XI(tp); x.width = x.height = 100
        mt.add_image(x, f'B{i}')
    else: print('no swatch', f_)
    mt.row_dimensions[i].height = 80

# ---------------------------------------------------------------- References
rf = wb.create_sheet('References')
rf.sheet_view.showGridLines = False
rf.column_dimensions['A'].width = 3
rf.cell(1, 2, 'Owner reference photos / Ảnh tham khảo của chủ đầu tư — style and quality we expect').font = font(bold=True, size=13, color=NAVY)
REFLIST = [('R12-v6/chair.jpg', 'Carved armchair (F01)'), ('R12-v6/cabinet.jpg', 'Jali cabinets (F08, F09)'), ('R12-v6/lanterns.jpg', 'Pierced brass lanterns (M02, M03)'),
           ('R12-v6/lanterns2.jpg', 'Lantern clusters (M01, M02)'), ('R12-v6/pendant.jpg', 'Pendant (M02)'), ('R12-v6/sconce.jpg', 'Sconce (M05)'),
           ('R12-v6/marblegrid.jpg', 'Floor grid with chip bands (C01)'), ('R12-v6/zellige.jpg', 'Zellige carpet (C02)'), ('R12-v6/border.jpg', 'Tile border (C04)')]
v8 = sorted(os.listdir(os.path.join(REFS, 'R12-v8')))
V8CAP = ['Owner reference (R12-v8)'] * len(v8)
REFLIST += [('R12-v8/' + f, c) for f, c in zip(v8, V8CAP)]
for i, (f_, cap) in enumerate(REFLIST):
    col, rr = 2 + (i % 3) * 5, 3 + (i // 3) * 17
    p = os.path.join(REFS, f_)
    im = PI.open(p).convert('RGB'); im.thumbnail((640, 480), PI.LANCZOS)
    tp = os.path.join(THUMB, 'ref_' + os.path.basename(f_).rsplit('.', 1)[0] + '.jpg'); im.save(tp, quality=85)
    x = XI(tp); sc = min(300 / im.size[0], 280 / im.size[1]); x.width, x.height = int(im.size[0] * sc), int(im.size[1] * sc)
    rf.add_image(x, rf.cell(rr, col).coordinate)
    rf.cell(rr + 15, col, cap).font = font(size=9, italic=True)
for j in range(2, 17): rf.column_dimensions[rf.cell(1, j).column_letter].width = 9

# ---------------------------------------------------------------- Design checks
dc = wb.create_sheet('Design checks')
dcols = (('No', 6), ('Topic', 14), ('Point to check / Điểm cần kiểm tra', 80), ('Who', 18), ('Affects', 16), ('Contractor answer / Trả lời nhà thầu', 44))
for j, (t, w) in enumerate(dcols, 1):
    c = dc.cell(1, j, t); c.font = font(bold=True, color='FFFFFF'); c.fill = HEAD; c.alignment = CEN; c.border = BOX
    dc.column_dimensions[c.column_letter].width = w
dc.freeze_panes = 'A2'
for i, q in enumerate(CHECKS, 2):
    for j, v in enumerate(q + (None,), 1):
        c = dc.cell(i, j, v); c.font = font(size=10, bold=(j == 1)); c.alignment = WRAP; c.border = BOX
    dc.cell(i, 6).fill = YEL; dc.row_dimensions[i].height = 44

wb.move_sheet('Design checks', offset=-(len(wb.sheetnames) - 3))
for s in wb.worksheets:
    s.page_setup.orientation = 'landscape'; s.page_setup.paperSize = 9; s.page_setup.fitToWidth = 1; s.page_setup.fitToHeight = 0; s.sheet_properties.pageSetUpPr.fitToPage = True
    s.sheet_properties.tabColor = {'Read me': GOLD, 'BOQ': NAVY, 'Summary': '2E7D32'}.get(s.title, '8C8C8C')
wb.save(OUT); print('saved', OUT, 'BOQ rows', nboq)
